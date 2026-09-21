import csv
import base64
import hashlib
import time
import os
import json
import math
import re
import uuid
from pathlib import Path

import httpx

from app.models.adapter import _strip_reasoning

# Removes the markers while keeping what is between them.
_THINK_TAGS = re.compile(r'</?think>', re.I)
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, text


SUPPORTED_EXTENSIONS = {'.txt', '.md', '.pdf', '.docx', '.pptx', '.csv', '.xlsx', '.xlsm', '.png', '.jpg', '.jpeg', '.tiff', '.bmp'}

# Uploads are stored on disk as "<uuid4>_<original filename>"; index and cite the
# original name so evidence lists read cleanly.
_UPLOAD_PREFIX = re.compile(r'^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}_')


def _display_name(path, metadata=None):
    return (metadata or {}).get('source_name') or _UPLOAD_PREFIX.sub('', Path(path).name)


def _decode_json(value, default):
    """Read a JSON column from either backend.

    psycopg hands back JSONB already decoded, SQLite hands back the TEXT it stored.
    Calling json.loads on the decoded dict raised "the JSON object must be str, bytes
    or bytearray, not dict" -- which only ever surfaced on PostgreSQL, and only on the
    synchronous search path, because the async one returns before reaching it.
    """
    if value is None:
        return default
    if isinstance(value, (dict, list)):
        return value
    return json.loads(value) if str(value).strip() else default


def _ranges(numbers):
    """Compact a page list for a human: [6,7,8,11] -> "6-8, 11"."""
    out, start, previous = [], None, None
    for number in sorted(numbers):
        if start is None:
            start = previous = number
        elif number == previous + 1:
            previous = number
        else:
            out.append(str(start) if start == previous else f'{start}-{previous}')
            start = previous = number
    if start is not None:
        out.append(str(start) if start == previous else f'{start}-{previous}')
    return ', '.join(out)


def _no_text_marker(name):
    """What to index when an image yields no readable text.

    Indexing the OCR noise instead was actively harmful: it is what retrieval
    returned as the file's evidence, and the model then reported the document as
    unreadable garbage rather than answering. A file with no extractable text is
    still answerable -- the orchestrator attaches the image itself to the task --
    so say that plainly and let the vision pass at question time do the work.
    """
    return (
        f'[{name}: image with no machine-readable text layer. Nothing could be transcribed by OCR. '
        'Any question about this file must be answered from the image itself, which is attached to '
        'the task, not from this entry.]'
    )


class RagService:
    """Local document index with durable chunks and Ollama embeddings."""

    def __init__(self, database_url='sqlite:///./orchestrator.db', ollama_base_url='http://localhost:11434', embedding_model='nomic-embed-text', vision_model='qwen2.5vl:3b', embedding_dimensions=768, keep_alive=None):
        self.engine = create_engine(database_url, future=True, pool_pre_ping=True)
        self.is_postgres = database_url.startswith('postgresql')
        self.ollama_url = ollama_base_url.rstrip('/') + '/api/embeddings'
        self.embedding_model = embedding_model
        self.embedding_dimensions = embedding_dimensions
        self.vision_url = ollama_base_url.rstrip('/') + '/api/chat'
        self.vision_model = vision_model
        # Sent on every embedding and vision call for the same reason as the chat
        # adapter: retrieval runs an embedding per query, so letting the embedding
        # model fall out of memory taxes every single search.
        self.keep_alive = keep_alive or os.getenv('OLLAMA_KEEP_ALIVE', '5m')
        # Tesseract is fast and offline but weak on scans, forms and handwriting.
        # Set OCR_PREFER_VISION=true to send images / scanned PDFs straight to the
        # local vision model instead (much better transcription, slower).
        self.prefer_vision = os.getenv('OCR_PREFER_VISION', 'false').lower() in {'1', 'true', 'yes'}
        # Ingest is synchronous -- an upload waits on this -- so the vision pass gets a
        # bounded budget rather than the caller's patience. A dense engineering drawing
        # measured at 393s on an 8B vision model here, so this will not always finish;
        # when it doesn't, ingest() falls back to an honest marker instead of noise.
        self.vision_timeout = int(os.getenv('OCR_VISION_TIMEOUT_SECONDS', '180'))
        # The budget above is for the WHOLE document, not per page. httpx applies a
        # timeout per request, so a per-page bound multiplied by the page count: a
        # twenty-page scan could hold an upload open for an hour. Pages are also
        # capped, because transcribing page 40 of a scan nobody will search is not
        # worth making the uploader wait. Measured here: ~51s for a typed page, ~111s
        # for a dense engineering drawing.
        self.max_vision_pages = int(os.getenv('OCR_MAX_VISION_PAGES', '5'))
        # Qwen3-VL is a thinking model and, measured here, ignores `think: false` on
        # this path -- it produced 2,937 characters of chain-of-thought anyway. It is
        # still sent for models that do honour it, but the transcription budget has to
        # assume thinking happens: capped at 1024 the reasoning consumed the entire
        # allowance and the call returned empty content. The budget below leaves room
        # for the answer after the thinking, and _vision_chat salvages the reasoning
        # text if a cap is still hit. Deliberately 4096 and not the 8192 the chat path
        # uses: measured on a dense drawing, 4096 completes a page in ~111s and 8192
        # pushes it past the 180s document budget, so the page times out and nothing is
        # indexed at all. A wider budget is not better when there is a deadline. The temperature is deliberately not lower: at 0.3
        # this model fell into a repetition loop on a dense drawing and never escaped.
        self._vision_options = {
            'think': False,
            'options': {'num_ctx': int(os.getenv('OLLAMA_NUM_CTX', '8192')), 'num_predict': 4096, 'temperature': 0.4},
        }
        metadata_type = 'JSONB' if self.is_postgres else 'TEXT'
        embedding_type = f'vector({embedding_dimensions})' if self.is_postgres else 'TEXT'
        if self.is_postgres:
            # Production Postgres is provisioned by the migrations in migrations/tier/;
            # just confirm the tables are present rather than issuing DDL at runtime.
            with self.engine.connect() as db:
                db.execute(text('SELECT 1 FROM rag_documents LIMIT 0'))
                db.execute(text('SELECT 1 FROM rag_chunks LIMIT 0'))
            return
        with self.engine.begin() as db:
            db.execute(text('''
                CREATE TABLE IF NOT EXISTS rag_documents (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, mime_type TEXT,
                    checksum TEXT NOT NULL, metadata ''' + metadata_type + ''' NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )'''))
            db.execute(text('''
                CREATE TABLE IF NOT EXISTS rag_chunks (
                    id TEXT PRIMARY KEY, document_id TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL, content TEXT NOT NULL,
                    embedding ''' + embedding_type + ''', metadata ''' + metadata_type + ''' NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES rag_documents(id)
                )'''))
            db.execute(text('CREATE INDEX IF NOT EXISTS idx_rag_chunks_document ON rag_chunks(document_id)'))

    def extract(self, path):
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix in {'.txt', '.md'}:
            return path.read_text(errors='ignore')
        if suffix == '.pdf':
            pages = []
            for index, page in enumerate(PdfReader(str(path)).pages, 1):
                try:
                    content = page.extract_text(extraction_mode='layout') or ''
                except TypeError:
                    content = page.extract_text() or ''
                pages.append(f'[Page {index}]\n{content}')
            text_layer = '\n'.join(pages)
            # An image-only / scanned PDF has almost no embedded text. Fall back to
            # rasterising each page and OCR'ing it (Tesseract; the async ingest path
            # additionally tries the local vision model).
            page_count = max(len(pages), 1)
            if len(re.sub(r'\s+', '', text_layer)) < 24 * page_count:
                ocr_layer = self._ocr_pdf(path)
                if len(re.sub(r'\s+', '', ocr_layer)) > len(re.sub(r'\s+', '', text_layer)):
                    return ocr_layer
            return text_layer
        if suffix == '.docx':
            return '\n'.join(p.text for p in Document(str(path)).paragraphs)
        if suffix == '.pptx':
            from pptx import Presentation
            slides = []
            for index, slide in enumerate(Presentation(str(path)).slides, 1):
                lines = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
                slides.append(f'[Slide {index}]\n' + '\n'.join(lines))
            return '\n\n'.join(slides)
        if suffix in {'.xlsx', '.xlsm'}:
            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f'[Sheet: {sheet.title}]')
                rows.extend(' | '.join('' if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
            return '\n'.join(rows)
        if suffix == '.csv':
            with path.open(newline='', encoding='utf-8-sig', errors='ignore') as stream:
                return '\n'.join(' | '.join(row) for row in csv.reader(stream))
        if suffix in {'.png', '.jpg', '.jpeg', '.tiff', '.bmp'}:
            return self._ocr(path)
        raise ValueError(f'UNSUPPORTED_DOCUMENT_TYPE: {suffix}')

    # Tesseract on a photo or an engineering drawing frequently returns neither an
    # error nor text -- it returns noise ('Y i , 7 "Y0DO us snd *juoaf') or a couple
    # of stray words. That was indexed as the document's entire content, so every
    # later answer about the image was reasoned from nonsense. Treat an unusable
    # pass the same as a failed one and let the local vision model transcribe it.
    _WORD = re.compile(r'[A-Za-z]{3,}')

    @classmethod
    def _ocr_is_unusable(cls, extracted):
        if not extracted or extracted.startswith('[OCR unavailable:'):
            return True
        if len(re.sub(r'\s+', '', extracted)) < 24:
            return True
        tokens = extracted.split()
        words = [t for t in tokens if cls._WORD.fullmatch(t.strip('.,:;!?()[]{}<>"\'`*|/\\-_=+$#@~^&%'))]
        return len(words) < max(3, len(tokens) * 0.4)

    @staticmethod
    def _ocr(path):
        try:
            import pytesseract
            from PIL import Image
            return pytesseract.image_to_string(Image.open(path))
        except Exception as exc:
            return f'[OCR unavailable: {exc}]'

    @staticmethod
    def _pdf_pages_without_text(path, min_chars=200):
        """``(page numbers whose text layer is too thin to be real text, page count)``.

        Judging the OCR *output* instead was unreliable in both directions: a page of
        tabular figures trips the noise heuristic, while several pages of Tesseract
        noise add up to something that looks substantial. The text layer is the thing
        actually being asked about, and reading it is nearly free.
        """
        import pymupdf
        document = pymupdf.open(str(path))
        try:
            return ([index + 1 for index in range(document.page_count)
                     if len(re.sub(r'\s+', '', document.load_page(index).get_text() or '')) < min_chars],
                    document.page_count)
        finally:
            document.close()

    @staticmethod
    def _pdf_page_images(path, dpi=200):
        """Yield (page_number, PNG bytes) for each page of a PDF, rendered locally."""
        import pymupdf
        document = pymupdf.open(str(path))
        try:
            for index in range(document.page_count):
                pixmap = document.load_page(index).get_pixmap(dpi=dpi)
                yield index + 1, pixmap.tobytes('png')
        finally:
            document.close()

    @classmethod
    def _ocr_pdf(cls, path):
        try:
            import io
            import pytesseract
            from PIL import Image
            pages = []
            for number, png in cls._pdf_page_images(path):
                text = pytesseract.image_to_string(Image.open(io.BytesIO(png)))
                pages.append(f'[Page {number}]\n{text}')
            return '\n'.join(pages)
        except Exception as exc:
            return f'[OCR unavailable: {exc}]'

    async def _vision_extract_pdf(self, path, only_pages=None):
        """Transcribe a PDF's picture pages within one bounded budget.

        Stops at max_vision_pages or when vision_timeout for the document is spent,
        whichever comes first, and records which pages were left untranscribed so the
        index does not imply the whole document was read.
        """
        try:
            import base64 as _b64
            wanted = set(only_pages) if only_pages is not None else None
            deadline = time.monotonic() + self.vision_timeout
            pages, skipped = [], []
            async with httpx.AsyncClient(timeout=self.vision_timeout) as client:
                for number, png in self._pdf_page_images(path):
                    if wanted is not None and number not in wanted:
                        continue
                    remaining = deadline - time.monotonic()
                    # 15s is below any observed single-page time, so starting a page
                    # with less than that left only guarantees a wasted timeout.
                    if len(pages) >= self.max_vision_pages or remaining < 15:
                        skipped.append(number)
                        continue
                    payload = {'model': self.vision_model, 'stream': False, 'keep_alive': self.keep_alive, **self._vision_options, 'messages': [{'role': 'user', 'content': 'Transcribe all visible text exactly, including handwritten text where legible. Return only the transcription.', 'images': [_b64.b64encode(png).decode('ascii')]}]}
                    pages.append(f'[Page {number}]\n' + await self._vision_chat(client, payload, timeout=remaining))
            if skipped:
                pages.append(f'[Pages {_ranges(skipped)} were not transcribed: the OCR budget '
                             f'({self.vision_timeout}s / {self.max_vision_pages} pages) was reached. '
                             'Ask about these pages directly -- they are sent to the vision model with the task.]')
            return '\n'.join(pages)
        except Exception as exc:
            return f'[OCR unavailable: install Tesseract or configure Ollama vision: {exc}]'

    @staticmethod
    def _chunks(text, size=1200, overlap=150):
        text = re.sub(r'\s+', ' ', text).strip()
        if not text:
            return []
        result = []
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            result.append(text[start:end])
            if end == len(text):
                break
            start = end - overlap
        return result

    async def _embed(self, text):
        try:
            async with httpx.AsyncClient(timeout=30) as client:
                response = await client.post(self.ollama_url, json={'model': self.embedding_model, 'prompt': text, 'keep_alive': self.keep_alive})
                response.raise_for_status()
                return response.json()['embedding']
        except Exception:
            return None

    async def ingest(self, path, metadata=None):
        path = Path(path)
        data = path.read_bytes()
        checksum = hashlib.sha256(data).hexdigest()
        metadata = metadata or {}
        with self.engine.connect() as db:
            existing = next((row for row in db.execute(text('SELECT id,metadata FROM rag_documents WHERE checksum=:checksum'), {'checksum': checksum}).fetchall() if (row.metadata if isinstance(row.metadata, dict) else json.loads(row.metadata or '{}')).get('tenant_id') == metadata.get('tenant_id')), None)
            if existing:
                return {'document_id': existing[0], 'name': _display_name(path, metadata), 'chunks': 0, 'existing': True}
        suffix = path.suffix.lower()
        is_image = suffix in {'.png', '.jpg', '.jpeg', '.tiff', '.bmp'}
        if is_image and self.prefer_vision:
            extracted_text = await self._vision_extract(path)
        elif suffix == '.pdf' and self.prefer_vision:
            extracted_text = await self._vision_extract_pdf(path)
        else:
            extracted_text = self.extract(path)
            if is_image and self._ocr_is_unusable(extracted_text):
                vision_text = await self._vision_extract(path)
                extracted_text = (
                    vision_text if not self._ocr_is_unusable(vision_text)
                    else _no_text_marker(_display_name(path, metadata))
                )
            elif suffix == '.pdf':
                # Which pages are pictures is decided by the text layer, not by how the
                # OCR output looks. The old test -- "[OCR unavailable:] or under 24
                # characters" -- passed 156 characters of Tesseract noise, so a scanned
                # drawing was indexed as gibberish and the vision model never ran on it.
                scanned, page_count = self._pdf_pages_without_text(path)
                if scanned:
                    vision_text = await self._vision_extract_pdf(path, only_pages=scanned)
                    if not self._ocr_is_unusable(vision_text):
                        # extract() only falls back to Tesseract when the WHOLE document is
                        # thin, so a mixed PDF still returns its real text layer and that is
                        # worth keeping beside the transcription. When every page is a
                        # picture, what extract() returned is the Tesseract noise this fix
                        # exists to get rid of -- replace it rather than concatenating it.
                        extracted_text = (
                            vision_text if len(scanned) == page_count
                            else extracted_text + '\n' + vision_text
                        )
        chunks = self._chunks(extracted_text)
        document_id = str(uuid.uuid4())
        vectors = [await self._embed(chunk) for chunk in chunks]
        if self.is_postgres and any(vector and len(vector) != self.embedding_dimensions for vector in vectors):
            raise ValueError(f'EMBEDDING_DIMENSION_MISMATCH: expected {self.embedding_dimensions} (set RAG_EMBEDDING_DIMENSIONS to match the model)')
        with self.engine.begin() as db:
            db.execute(text('INSERT INTO rag_documents(id,name,mime_type,checksum,metadata) VALUES(:id,:name,:mime,:checksum,' + ('CAST(:metadata AS JSONB)' if self.is_postgres else ':metadata') + ')'), {'id': document_id, 'name': _display_name(path, metadata), 'mime': metadata.get('mime_type'), 'checksum': checksum, 'metadata': json.dumps(metadata)})
            for index, (chunk, vector) in enumerate(zip(chunks, vectors)):
                embedding = json.dumps(vector) if vector else None
                db.execute(text('INSERT INTO rag_chunks(id,document_id,chunk_index,content,embedding,metadata) VALUES(:id,:document,:index,:content,' + ('CAST(:embedding AS vector)' if self.is_postgres else ':embedding') + ',' + ('CAST(:metadata AS JSONB)' if self.is_postgres else ':metadata') + ')'), {'id': str(uuid.uuid4()), 'document': document_id, 'index': index, 'content': chunk, 'embedding': embedding, 'metadata': json.dumps(metadata)})
        return {'document_id': document_id, 'name': _display_name(path, metadata), 'chunks': len(chunks), 'embedded': sum(vector is not None for vector in vectors)}

    def ingest_sync(self, path, metadata=None):
        """Index a document without network access for synchronous tools."""
        path = Path(path)
        data = path.read_bytes()
        checksum = hashlib.sha256(data).hexdigest()
        metadata = metadata or {}
        with self.engine.connect() as db:
            existing = next((row for row in db.execute(text('SELECT id,metadata FROM rag_documents WHERE checksum=:checksum'), {'checksum': checksum}).fetchall() if (row.metadata if isinstance(row.metadata, dict) else json.loads(row.metadata or '{}')).get('tenant_id') == metadata.get('tenant_id')), None)
            if existing:
                return {'document_id': existing[0], 'name': _display_name(path, metadata), 'chunks': 0, 'embedded': 0, 'existing': True}
        extracted_text = self.extract(path)
        chunks = self._chunks(extracted_text)
        document_id = str(uuid.uuid4())
        with self.engine.begin() as db:
            db.execute(text('INSERT INTO rag_documents(id,name,mime_type,checksum,metadata) VALUES(:id,:name,:mime,:checksum,' + ('CAST(:metadata AS JSONB)' if self.is_postgres else ':metadata') + ')'), {'id': document_id, 'name': _display_name(path, metadata), 'mime': metadata.get('mime_type'), 'checksum': checksum, 'metadata': json.dumps(metadata)})
            for index, chunk in enumerate(chunks):
                db.execute(text('INSERT INTO rag_chunks(id,document_id,chunk_index,content,embedding,metadata) VALUES(:id,:document,:index,:content,NULL,' + ('CAST(:metadata AS JSONB)' if self.is_postgres else ':metadata') + ')'), {'id': str(uuid.uuid4()), 'document': document_id, 'index': index, 'content': chunk, 'metadata': json.dumps(metadata)})
        return {'document_id': document_id, 'name': _display_name(path, metadata), 'chunks': len(chunks), 'embedded': 0}

    async def _vision_chat(self, client, payload, timeout=None):
        """POST one vision request, retrying without `think` for models that reject it."""
        extra = {'timeout': timeout} if timeout else {}
        response = await client.post(self.vision_url, json=payload, **extra)
        if response.status_code == 400 and 'think' in response.text.lower():
            payload = {k: v for k, v in payload.items() if k != 'think'}
            response = await client.post(self.vision_url, json=payload, **extra)
        response.raise_for_status()
        message = response.json().get('message', {})
        content = (message.get('content') or '').strip()
        # A thinking model that exhausts num_predict before it starts answering returns
        # empty content with the reasoning intact. That reasoning contains the text it
        # read off the page, so it beats indexing nothing. Same salvage as OllamaAdapter.
        # Qwen3-VL also ignores `think: false` and writes <think> blocks straight into
        # content, and indexing a transcription wrapped in the model's own deliberation
        # pollutes every later search -- so prefer the answer with the reasoning removed.
        #
        # But on a dense page this model frequently spends the whole budget reasoning and
        # never writes an answer at all (measured: 8,491 characters of thinking, zero of
        # content), and the transcription it read off the page is in that reasoning. When
        # stripping leaves nothing worth indexing, keep the reasoning and drop only the
        # tags: imperfect text that contains the drawing's callouts beats clean text that
        # contains nothing.
        raw = content or (message.get('thinking') or '')
        cleaned = _strip_reasoning(raw)
        return cleaned if not self._ocr_is_unusable(cleaned) else _THINK_TAGS.sub('', raw).strip()

    async def _vision_extract(self, path):
        try:
            encoded = base64.b64encode(path.read_bytes()).decode('ascii')
            payload = {'model': self.vision_model, 'stream': False, 'keep_alive': self.keep_alive, **self._vision_options, 'messages': [{'role': 'user', 'content': 'Transcribe all visible text exactly. Include handwritten text where legible. Return only the transcription.', 'images': [encoded]}]}
            async with httpx.AsyncClient(timeout=self.vision_timeout) as client:
                return await self._vision_chat(client, payload)
        except Exception as exc:
            return f'[OCR unavailable: install Tesseract or configure Ollama vision: {exc}]'

    async def search(self, query, top_k=5, metadata=None, file_ids=None):
        query_vector = await self._embed(query)
        return self._search_rows(query, query_vector, top_k, metadata, file_ids)

    def search_sync(self, query, top_k=5, metadata=None, file_ids=None):
        """Search without network access for the synchronous tool dispatcher."""
        return self._search_rows(query, None, top_k, metadata, file_ids)

    def _search_rows(self, query, query_vector, top_k, metadata, file_ids=None):
        with self.engine.connect() as db:
            if self.is_postgres and query_vector:
                filters = []
                params = {'embedding': json.dumps(query_vector), 'limit': min(max(top_k * 4, top_k, 1), 50)}
                if metadata:
                    filters.append('c.metadata @> CAST(:metadata AS JSONB)')
                    params['metadata'] = json.dumps(metadata)
                if file_ids is not None:
                    if not file_ids:
                        return []
                    placeholders = []
                    for index, file_id in enumerate(file_ids):
                        key = f'file_id_{index}'
                        placeholders.append(f':{key}')
                        params[key] = file_id
                    filters.append("c.metadata->>'file_id' IN (" + ','.join(placeholders) + ')')
                where = (' WHERE ' + ' AND '.join(filters)) if filters else ''
                rows = db.execute(text('SELECT c.id,c.document_id,c.content,c.embedding,c.metadata,d.name, 1 - (c.embedding <=> CAST(:embedding AS vector)) AS score FROM rag_chunks c JOIN rag_documents d ON d.id=c.document_id' + where + ' ORDER BY c.embedding <=> CAST(:embedding AS vector) LIMIT :limit'), params).fetchall()
                candidates = [self._hit(row, float(row.score or 0), row.embedding) for row in rows]
                return self._rerank(query, candidates, top_k)
            rows = db.execute(text('SELECT c.id,c.document_id,c.content,c.embedding,c.metadata,d.name FROM rag_chunks c JOIN rag_documents d ON d.id=c.document_id')).fetchall()
        scored = []
        query_words = set(re.findall(r'\w+', query.lower()))
        for chunk_id, document_id, content, embedding, chunk_metadata, name in rows:
            item_metadata = _decode_json(chunk_metadata, {})
            if metadata and any(item_metadata.get(key) != value for key, value in metadata.items()):
                continue
            if file_ids is not None and item_metadata.get('file_id') not in file_ids:
                continue
            score = self._cosine(query_vector, _decode_json(embedding, [])) if query_vector and embedding else self._lexical(query_words, content)
            scored.append({'chunk_id': chunk_id, 'document_id': document_id, 'source': name, 'content': content, 'score': round(score, 6), 'metadata': item_metadata, 'retrieval_method': 'embedding' if query_vector and embedding else 'lexical'})
        return self._rerank(query, scored, top_k)

    def _rerank(self, query, candidates, top_k):
        query_words = set(re.findall(r'\w+', query.lower()))
        for item in candidates:
            lexical = self._lexical(query_words, item['content'])
            item['score'] = round((0.75 * item['score']) + (0.25 * lexical), 6)
            item['retrieval_method'] = item.get('retrieval_method', 'lexical') + '+local_rerank'
        return sorted(candidates, key=lambda item: item['score'], reverse=True)[:top_k]

    async def evaluate(self, cases, metadata=None, file_ids=None):
        results = []
        reciprocal_ranks = []
        hits = 0
        for case in cases:
            expected = set(case.get('expected_file_ids') or [])
            found = await self.search(case['query'], case.get('top_k', 5), metadata, file_ids)
            found_ids = [item['metadata'].get('file_id') for item in found]
            rank = next((index + 1 for index, file_id in enumerate(found_ids) if file_id in expected), None)
            if rank:
                hits += 1
                reciprocal_ranks.append(1 / rank)
            else:
                reciprocal_ranks.append(0)
            results.append({'query': case['query'], 'expected_file_ids': list(expected), 'found_file_ids': found_ids, 'hit': bool(rank), 'rank': rank})
        total = len(cases)
        return {'cases': results, 'metrics': {'count': total, 'hit_rate': hits / total if total else 0, 'mrr': sum(reciprocal_ranks) / total if total else 0}}

    @staticmethod
    def _hit(row, score, embedding):
        metadata = row.metadata if isinstance(row.metadata, dict) else json.loads(row.metadata or '{}')
        return {'chunk_id': row.id, 'document_id': row.document_id, 'source': row.name, 'content': row.content, 'score': round(score, 6), 'metadata': metadata, 'retrieval_method': 'pgvector'}

    def delete_document(self, file_id):
        with self.engine.begin() as db:
            document_ids = db.execute(text('SELECT id FROM rag_documents WHERE ' + ('metadata @> CAST(:metadata AS JSONB)' if self.is_postgres else 'metadata LIKE :metadata')), {'metadata': json.dumps({'file_id': file_id}) if self.is_postgres else '%"file_id": "' + file_id + '"%'}).fetchall()
            for (document_id,) in document_ids:
                db.execute(text('DELETE FROM rag_chunks WHERE document_id=:id'), {'id': document_id})
                db.execute(text('DELETE FROM rag_documents WHERE id=:id'), {'id': document_id})
        return len(document_ids)

    @staticmethod
    def _lexical(words, content):
        tokens = set(re.findall(r'\w+', content.lower()))
        return len(words & tokens) / max(len(words), 1)

    @staticmethod
    def _cosine(left, right):
        if not left or not right or len(left) != len(right):
            return 0.0
        dot = sum(a * b for a, b in zip(left, right))
        norm = math.sqrt(sum(a * a for a in left)) * math.sqrt(sum(b * b for b in right))
        return dot / norm if norm else 0.0